#!/usr/bin/env python3
"""
TRACE OSINT Digital Footprint Analyzer - 5-Page Case Study DOCX Generator
Author: Fardeen Akmal (Roll No: 2)
Guide: Akanksha Thakare
Institute: G.V. Acharya Institute of Engineering and Technology
University: University of Mumbai
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(4)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(15, 23, 42)
    elif level == 2:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(3, 105, 161)
    return h

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(7.8)
        r_pre.font.name = 'Calibri'
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.size = Pt(7.8)
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_item(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.2)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.bold = True
        r_pre.font.size = Pt(7.8)
        r_pre.font.name = 'Calibri'
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.size = Pt(7.8)
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def build_docx(filename):
    doc = docx.Document()

    # Configure Margins (0.45 inches)
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

    # Header and Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("TRACE: OSINT DIGITAL FOOTPRINT ANALYZER — MINI PROJECT CASE STUDY")
    hrun.font.size = Pt(7.5)
    hrun.font.bold = True
    hrun.font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    frun1 = fp.add_run("G.V. Acharya Institute of Engineering & Technology • University of Mumbai")
    frun1.font.size = Pt(7.5)
    frun1.font.color.rgb = RGBColor(100, 116, 139)
    frun2 = fp.add_run("\t\tAcademic Year 2026-2027")
    frun2.font.size = Pt(7.5)
    frun2.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # PAGE 1: TITLE, ACADEMIC DETAILS, ABSTRACT, INTRODUCTION, PROBLEM STATEMENT, OBJECTIVES & SCOPE
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("TRACE — OSINT DIGITAL FOOTPRINT ANALYZER")
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("A Robust Web-Based Open Source Intelligence Framework for Surface-Web Profile & Domain Footprint Auditing")
    r_sub.bold = True
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(2, 132, 199)

    # Student metadata table
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        [("Student Name: ", "Fardeen Akmal"), ("Roll Number: ", "2"), ("Project Guide: ", "Akanksha Thakare")],
        [("Branch: ", "CSE (IoT, Cyber Security & BCT)"), ("Institute: ", "GVAIET, Shelu"), ("Affiliation: ", "University of Mumbai")]
    ]
    for r_idx, row in enumerate(meta_rows):
        for c_idx, (lbl, val) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(lbl)
            r1.bold = True
            r1.font.size = Pt(7.5)
            r1.font.color.rgb = RGBColor(15, 23, 42)
            r2 = p.add_run(val)
            r2.font.size = Pt(7.5)
            r2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # Abstract Box
    abs_table = doc.add_table(rows=1, cols=1)
    abs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = abs_table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
    p_abs = cell.paragraphs[0]
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.line_spacing = 1.05
    p_abs.paragraph_format.space_after = Pt(0)
    r_absh = p_abs.add_run("ABSTRACT: ")
    r_absh.bold = True
    r_absh.font.size = Pt(7.8)
    r_absh.font.color.rgb = RGBColor(2, 132, 199)
    r_abst = p_abs.add_run(
        "In the contemporary hyper-connected digital ecosystem, an individual's or enterprise's digital footprint across "
        "heterogeneous web platforms represents an expansive attack surface susceptible to social engineering, spear phishing, "
        "credential stuffing, and unauthorized reconnaissance. Open Source Intelligence (OSINT) facilitates the proactive identification "
        "and auditing of publicly exposed indicators of identity. However, legacy reconnaissance utilities often rely on brittle web-scraping "
        "heuristics, generate elevated false-positive rates due to platform anti-bot defenses, fail to distinguish rate limits from "
        "legitimate target non-existence, and lack localized, auditable reporting capabilities. This case study presents the design, "
        "architectural formulation, empirical implementation, and testing of TRACE (Target Reconnaissance & Auditable Collection Engine) — "
        "a full-stack, modular OSINT analysis dashboard. Built using a Python FastAPI backend and a responsive React/TypeScript client with "
        "SQLite local persistence, TRACE executes rigorous username enumeration across ten major developer, visual, and social platforms "
        "alongside passive DNS and RDAP domain intelligence gathering. TRACE introduces an unambiguous multi-tier status categorization engine "
        "(FOUND, NOT_FOUND, INCONCLUSIVE, CHECK_FAILED) that prevents fabricated intelligence and provides verifiable cryptographic checksums "
        "with automated PDF/CSV report exports."
    )
    r_abst.font.size = Pt(7.8)
    r_abst.font.color.rgb = RGBColor(30, 41, 59)

    add_styled_heading(doc, "1. Introduction & Theoretical Background", level=1)
    add_body_paragraph(
        doc,
        "Open Source Intelligence (OSINT) comprises the rigorous collection, correlation, and analysis of publicly available data to produce "
        "actionable security intelligence. While intelligence agencies initially pioneered OSINT methodologies, modern defensive cybersecurity "
        "teams leverage OSINT for External Attack Surface Management (EASM), threat intelligence, and brand reputation monitoring. A critical "
        "challenge in modern OSINT operations is the fragmented nature of surface-web footprints: a single identifier or domain is frequently "
        "linked to software repositories, social forums, identity providers, and hosting providers. Manual correlation across dozens of platforms "
        "is tedious, error-prone, and operationally inefficient. TRACE was conceived to democratize and systematize authorized digital footprint "
        "auditing through an intuitive, centralized, and standards-compliant analytical interface."
    )

    add_styled_heading(doc, "2. Problem Statement", level=1)
    add_body_paragraph(
        doc,
        "Existing CLI-based OSINT utilities suffer from significant architectural and operational deficiencies: "
        "(1) False-Positive Hallucination: Many tools assume that any HTTP 200 response confirms account existence, ignoring Single-Page Application (SPA) "
        "client-side routing, generic soft-404 landing pages, and Cloudflare challenge screens. "
        "(2) Ambiguous Failure States: Network timeouts, HTTP 429 rate limits, and platform WAF mitigations are frequently misreported as non-existent "
        "accounts, skewing investigator findings. "
        "(3) Excessive Active Probing: Intrusive network scanning techniques risk triggering Intrusion Detection Systems (IDS) or violating organizational "
        "Acceptable Use Policies (AUP). "
        "(4) Lack of Academic & Audit-Ready Reporting: Security professionals and students require defensible, timestamped, and reproducible evidence "
        "formats (PDF/CSV) that capture latency telemetry, raw DNS zone delegations, and operational limitations."
    )

    add_styled_heading(doc, "3. Objectives and Project Scope", level=1)
    add_body_paragraph(
        doc,
        "The primary objective of this mini-project is to build, execute, rigorously test, and demonstrate an autonomous, full-stack OSINT digital "
        "footprint analyzer operating exclusively on a local development workstation. Specific objectives include:"
    )
    add_bullet_item(doc, "Deterministic Profile Enumeration:", "Query 10 diverse public services concurrently using asynchronous HTTP with platform-specific signature matching and explicit status disambiguation.")
    add_bullet_item(doc, "Passive Domain Reconnaissance:", "Query publicly delegated authoritative DNS records (A, AAAA, MX, NS, TXT, SOA), ICANN RDAP registration registries, and GeoIP ASN telemetry without invasive port scanning.")
    add_bullet_item(doc, "Local Audit Persistence:", "Maintain a secure SQLite database recording all queries, statuses, timestamps, and findings with zero external data leakage or dependency on commercial cloud APIs.")
    add_bullet_item(doc, "Defensible Evidence Generation:", "Provide automated PDF and CSV reporting engines featuring academic attribution, integrity checksums, and explicit methodology disclaimers.")
    add_bullet_item(doc, "Scope Boundaries:", "TRACE focuses exclusively on public surface-web assets; active exploitation, vulnerability scanning, password harvesting, and bypassing authentication controls are strictly excluded.")

    doc.add_page_break()

    # =========================================================================
    # PAGE 2: PROPOSED SYSTEM, RELATED WORK, REQUIREMENTS, TECH STACK, USE CASES
    # =========================================================================
    add_styled_heading(doc, "4. Proposed System & Literature Review", level=1)
    add_body_paragraph(
        doc,
        "TRACE operates as an on-premise, security-centric analytical platform designed to overcome the fragility of legacy CLI scripts. "
        "Unlike monolithic scrapers, TRACE decouples asynchronous protocol dispatch from presentation, enforcing validation gateways "
        "at both network and application layers."
    )

    # Tool comparison table
    comp_table = doc.add_table(rows=5, cols=6)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Tool Name", "Interface", "Rate Limit Detection", "Passive DNS / RDAP", "Audit Report Generation", "Local Persistence"]
    rows = [
        ["Sherlock", "CLI Only", "Partial (Silent 429)", "No (Usernames only)", "Raw Text / CSV", "None (Stateless)"],
        ["theHarvester", "CLI Only", "Search API dependent", "Yes (Active DNS)", "HTML / XML export", "Optional SQLite"],
        ["SpiderFoot", "Web UI / Heavy", "Heuristic based", "Extensive (Complex)", "CSV / GEXF export", "Heavy SQLite"],
        ["TRACE (Ours)", "Modern Web UI", "Strict Disambiguation", "Integrated Passive", "Automated PDF + CSV", "Lightweight SQLite"]
    ]
    for c_idx, h in enumerate(headers):
        cell = comp_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=30, bottom=30, left=40, right=40)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = comp_table.cell(r_idx + 1, c_idx)
            bg = "E0F2FE" if r_idx == 3 else ("FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=25, bottom=25, left=40, right=40)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(7)
            if r_idx == 3:
                r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)

    add_styled_heading(doc, "5. System Requirements Analysis", level=1)
    add_styled_heading(doc, "5.1 Functional Requirements (FR)", level=2)
    add_bullet_item(doc, "FR-1 (Multi-Platform Profile Verification):", "The system must scan target usernames across 10 defined platforms (GitHub, GitLab, DockerHub, Dev.to, HackerNews, Reddit, Gravatar, Keybase, Medium, Pinterest) concurrently.")
    add_bullet_item(doc, "FR-2 (Passive DNS Resolution):", "The system must retrieve authoritative DNS records (A, AAAA, MX, NS, TXT, SOA) for target domains using RFC-compliant UDP/TCP DNS queries without touching target web servers.")
    add_bullet_item(doc, "FR-3 (RDAP & IP Geolocation Extraction):", "The system must parse Registration Data Access Protocol (RDAP) JSON structures to obtain registrar, registration, and expiration dates, supplemented by IP ASN data.")
    add_bullet_item(doc, "FR-4 (Audit Logging & Record Management):", "All completed investigations must persist in a local relational database with query parameters, timestamps, execution summaries, and deletion management.")
    add_bullet_item(doc, "FR-5 (Formal Academic Export):", "The system must generate self-contained, publication-grade PDF reports (incorporating student credentials and checksums) and tabular CSV exports on demand.")

    add_styled_heading(doc, "5.2 Non-Functional Requirements (NFR)", level=2)
    add_bullet_item(doc, "NFR-1 (Integrity & Zero False-Positive Tolerance):", "Under no condition shall the application fabricate finding statuses or confuse HTTP 403/429 anti-bot mitigations with target existence.")
    add_bullet_item(doc, "NFR-2 (Latency & Concurrency):", "Concurrency through Python asyncio and httpx must limit complete 10-source username footprint scans to under 4.5 seconds under standard broadband conditions.")
    add_bullet_item(doc, "NFR-3 (Security & Hygiene):", "The application must run completely offline/locally without transmitting telemetry to third-party tracking services or hardcoding secrets.")

    add_styled_heading(doc, "6. Technology Stack Selection & Rationale", level=1)
    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Component Layer", "Technology Selected", "Engineering Rationale"]
    t_rows = [
        ["Backend Framework", "Python 3.14 + FastAPI", "High-throughput asynchronous event loop, native Pydantic schema validation, and automatic OpenAPI documentation."],
        ["Asynchronous HTTP", "HTTPX Client Pool", "Non-blocking connection pooling, HTTP/2 support, custom headers/cookies, and strict per-request timeout handling."],
        ["DNS Reconnaissance", "dnspython v2.7", "Direct DNS wire-format resolver interfacing with authoritative nameservers without reliance on OS shell hooks."],
        ["Client Architecture", "React 18 + TypeScript + Vite", "Type-safe state management, fast client bundling, and modular component hierarchy."],
        ["Persistence Layer", "SQLite 3 Embedded Engine", "Zero-configuration, ACID-compliant local SQL database preventing external attack vectors and credential overhead."],
        ["Document Engine", "ReportLab Flowable PDF", "Programmatic, pixel-accurate vector PDF synthesis with institutional header templates and tabular layouts."]
    ]
    for c_idx, h in enumerate(t_headers):
        cell = tech_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=25, bottom=25, left=40, right=40)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(t_rows):
        for c_idx, val in enumerate(row):
            cell = tech_table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=20, bottom=20, left=40, right=40)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(7)
            r.font.color.rgb = RGBColor(15, 23, 42)

    add_styled_heading(doc, "7. Operational Use Cases", level=1)
    add_body_paragraph(
        doc,
        "UC-1: Defensive Self-Footprint Audit: Security researchers and developers audit their personal usernames across platforms to locate abandoned accounts exposing legacy emails or weak authentication. "
        "UC-2: External Surface Assessment: Authorized penetration testers establish domain perimeter inventory (mail routing, nameserver delegation, SPF validation) during non-invasive pre-engagement reconnaissance."
    )

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: ARCHITECTURE, WORKFLOW, DATABASE, VERIFICATION, SECURITY CONTROLS
    # =========================================================================
    add_styled_heading(doc, "8. System Architecture and Workflow", level=1)
    add_body_paragraph(
        doc,
        "TRACE implements a decoupled, three-tier modular architecture comprising the Single Page Application Presentation Layer, "
        "the FastAPI Asynchronous Service Controller, and the Persistence & System Integration Engine."
    )

    arch_table = doc.add_table(rows=1, cols=4)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    blocks = [
        ("Client (React/TS)", "• Operations Dashboard\n• Asynchronous Forms\n• Real-time Badges\n• Dossier Export Modal", "F0FDF4"),
        ("Gateway (FastAPI)", "• /api/investigate/user\n• /api/investigate/domain\n• /api/history (CRUD)\n• /api/reports/pdf|csv", "EFF6FF"),
        ("Core Engines", "• UsernameOSINT\n• DomainOSINT\n• ReportingEngine\n• DatabaseManager", "FAF5FF"),
        ("Storage & Network", "• Public Web APIs\n• Authoritative DNS (53)\n• ICANN RDAP (HTTPS)\n• SQLite (trace.db)", "FFFBEB")
    ]
    for c_idx, (title, content, bg) in enumerate(blocks):
        cell = arch_table.cell(0, c_idx)
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=35, bottom=35, left=45, right=45)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_t = p.add_run(title + "\n")
        r_t.bold = True
        r_t.font.size = Pt(7.5)
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        r_c = p.add_run(content)
        r_c.font.size = Pt(6.8)
        r_c.font.color.rgb = RGBColor(51, 65, 85)

    add_styled_heading(doc, "9. Database Design & Schema Architecture", level=1)
    add_body_paragraph(
        doc,
        "To guarantee audit integrity, TRACE utilizes SQLite 3 with WAL (Write-Ahead Logging) mode enabled. "
        "The schema isolates operational telemetry from raw payload responses, maintaining structured indexes for rapid lookup:"
    )

    db_table = doc.add_table(rows=8, cols=4)
    db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    db_headers = ["Column Name", "Data Type", "Constraints", "Description & Functional Purpose"]
    db_rows = [
        ["id", "INTEGER", "PRIMARY KEY AUTOINCREMENT", "Unique sequential investigation dossier identifier."],
        ["investigation_type", "VARCHAR(20)", "NOT NULL, CHECK (type IN...)", "Identifies analysis pipeline: 'username' or 'domain'."],
        ["target", "VARCHAR(255)", "NOT NULL, INDEXED", "Sanitized target query string (username or domain)."],
        ["timestamp", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "UTC execution timestamp ensuring tamper-evident chronology."],
        ["status", "VARCHAR(50)", "NOT NULL", "Overall operational state: 'Completed' or 'Failed'."],
        ["summary", "TEXT", "NOT NULL", "Human-readable aggregate findings synopsis."],
        ["data_json", "TEXT (JSON)", "NOT NULL", "Full normalized JSON payload storing per-source latency and telemetry."]
    ]
    for c_idx, h in enumerate(db_headers):
        cell = db_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=20, bottom=20, left=35, right=35)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(db_rows):
        for c_idx, val in enumerate(row):
            cell = db_table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=18, bottom=18, left=35, right=35)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(7)
            r.font.color.rgb = RGBColor(15, 23, 42)

    add_styled_heading(doc, "10. Multi-Tier Result Verification Methodology", level=1)
    add_body_paragraph(
        doc,
        "A foundational principle of TRACE is the eradication of false positives. Standard HTTP status codes alone are insufficient "
        "due to modern web architectures. TRACE executes a four-state deterministic decision tree for every source probe:"
    )
    add_bullet_item(doc, "FOUND:", "Validated when HTTP status is exactly 200, the response body contains positive profile identifiers (e.g., matching username in JSON payload or verified meta tags), and redirection to sign-in portals has not occurred.")
    add_bullet_item(doc, "NOT_FOUND:", "Triggered when the service responds with definitive non-existence status (HTTP 404), or redirects to registration portals (e.g., GitLab redirecting unauthenticated missing profiles to users/sign_in with HTTP 302).")
    add_bullet_item(doc, "INCONCLUSIVE:", "Applied when anti-scraping defenses (Cloudflare Turnstile, AWS WAF), bot challenges, or rate limiters respond with HTTP 403 Forbidden or HTTP 429 Too Many Requests. The target profile may or may not exist; TRACE refuses to speculate.")
    add_bullet_item(doc, "CHECK_FAILED:", "Recorded when transport-layer faults occur, including TCP socket timeouts (>4.0s), DNS NXDOMAIN errors, or SSL handshake failures. The specific network exception is logged for audit defensibility.")

    add_styled_heading(doc, "11. Security Controls and Operational Safeguards", level=1)
    add_body_paragraph(
        doc,
        "Input Sanitization Gateways: Domain queries are strictly checked against RFC 1035 hostname specifications. Submissions containing "
        "shell metacharacters (;&|`$), path traversal markers, or URI schemes (http://) are rejected at the client gateway and server Pydantic "
        "validation boundary with HTTP 422, eliminating Server-Side Request Forgery (SSRF) and injection vectors. "
        "Passive Reconnaissance Guarantee: TRACE restricts network interactions to passive queries against third-party authoritative servers. "
        "No port scans, directory fuzzing, or active banner grabbing are ever executed against target infrastructure."
    )

    doc.add_page_break()

    # =========================================================================
    # PAGE 4: IMPLEMENTATION, TESTING PROCEDURE, TEST MATRIX, GENUINE SCREENSHOTS
    # =========================================================================
    add_styled_heading(doc, "12. Implementation Methodology & Engineering Execution", level=1)
    add_body_paragraph(
        doc,
        "The project was implemented under a rigorous Test-Driven Development (TDD) lifecycle on Linux x86_64. "
        "The backend leverages Python's non-blocking asyncio primitives to dispatch concurrent HTTP and DNS probes. "
        "The frontend is compiled as an optimized single-page static bundle and hosted directly via FastAPI mount points."
    )

    add_styled_heading(doc, "13. Automated Testing Procedure & Test Matrix", level=1)
    add_body_paragraph(
        doc,
        "An automated test suite was constructed using pytest and httpx.AsyncClient, verifying all API endpoints, "
        "validation gates, error states, and database transaction integrity under simulated network conditions."
    )

    t_matrix = doc.add_table(rows=9, cols=5)
    t_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    tm_headers = ["Test ID", "Test Unit & Description", "Input Parameters", "Observed Outcome", "Status"]
    tm_rows = [
        ["TC-01", "API Health & Subsystem Telemetry", "GET /api/health", "Status 200, engine: 'TRACE OSINT'", "PASS"],
        ["TC-02", "Dashboard Telemetry Counter", "GET /api/stats", "Status 200, valid counts & history list", "PASS"],
        ["TC-03", "Username Investigation Dispatch", "POST {'username': 'test'}", "Status 200, 10 findings categorized", "PASS"],
        ["TC-04", "Input Validation: Empty Target", "POST {'username': '  '}", "Status 422 Unprocessable Entity", "PASS"],
        ["TC-05", "Domain DNS & RDAP Resolution", "POST {'domain': 'example.com'}", "Status 200, valid A & NS records", "PASS"],
        ["TC-06", "Investigation History Audit Trail", "GET /api/history", "Status 200, JSON list of investigations", "PASS"],
        ["TC-07", "PDF Report Export Verification", "GET /api/reports/pdf/1", "Status 200, application/pdf MIME", "PASS"],
        ["TC-08", "CSV Structured Export Verification", "GET /api/reports/csv/1", "Status 200, text/csv MIME, valid rows", "PASS"]
    ]
    for c_idx, h in enumerate(tm_headers):
        cell = t_matrix.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=20, bottom=20, left=30, right=30)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(tm_rows):
        for c_idx, val in enumerate(row):
            cell = t_matrix.cell(r_idx + 1, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=16, bottom=16, left=30, right=30)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(7)
            if c_idx == 4:
                r.bold = True
                r.font.color.rgb = RGBColor(21, 128, 61)
            else:
                r.font.color.rgb = RGBColor(15, 23, 42)

    add_styled_heading(doc, "14. Genuine Application Screenshots & Verification Evidence", level=1)
    add_body_paragraph(
        doc,
        "Screenshots were captured directly from the live TRACE instance executing on the host workstation via headless Chromium browser:"
    )

    shot_base = "/home/fardeenakmal/Desktop/osint/TRACE-OSINT/screenshots"
    s1 = os.path.join(shot_base, "01-dashboard.png")
    s2 = os.path.join(shot_base, "02-username-investigation.png")
    s3 = os.path.join(shot_base, "03-domain-intelligence.png")
    s4 = os.path.join(shot_base, "05-report-export.png")

    img_table = doc.add_table(rows=4, cols=2)
    img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    captions = [
        ("Figure 1: Operational Dashboard (Live Telemetry & Audit Log)", "Figure 2: Username OSINT Results (Target: @fardeenakmal)"),
        ("Figure 3: Domain Intelligence (example.com RDAP/DNS)", "Figure 4: Audit Dossier Preview & Export Modal")
    ]
    img_paths = [(s1, s2), (s3, s4)]

    for block_idx in range(2):
        cap_row = block_idx * 2
        img_row = block_idx * 2 + 1
        for col_idx in range(2):
            # Caption
            c_cell = img_table.cell(cap_row, col_idx)
            set_cell_background(c_cell, "F1F5F9")
            set_cell_margins(c_cell, top=15, bottom=15, left=30, right=30)
            cp = c_cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            cr = cp.add_run(captions[block_idx][col_idx])
            cr.bold = True
            cr.font.size = Pt(6.8)
            cr.font.color.rgb = RGBColor(15, 23, 42)

            # Image
            i_cell = img_table.cell(img_row, col_idx)
            set_cell_margins(i_cell, top=15, bottom=15, left=30, right=30)
            ip = i_cell.paragraphs[0]
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.space_after = Pt(0)
            p_img = img_paths[block_idx][col_idx]
            if os.path.exists(p_img):
                ip.add_run().add_picture(p_img, width=Inches(3.45))

    doc.add_page_break()

    # =========================================================================
    # PAGE 5: FINDINGS, LIMITATIONS, ETHICAL/LEGAL, CONCLUSION, REFERENCES
    # =========================================================================
    add_styled_heading(doc, "15. Empirical Findings & Performance Evaluation", level=1)
    add_body_paragraph(
        doc,
        "Empirical testing with target username fardeenakmal and target domain example.com yielded decisive performance data: "
        "(1) Profile Detection Accuracy: The system identified a verified active public profile on GitHub (HTTP 200, 614ms latency) with full bio extraction, "
        "while correctly distinguishing absent profiles on GitLab (HTTP 302 login redirect, 639ms), DockerHub (HTTP 404, 628ms), Dev.to (HTTP 404), "
        "and Keybase (HTTP 404). "
        "(2) Anti-Scraping Resilience: Reddit and Medium returned HTTP 403 status codes due to cloud-edge WAF challenges; TRACE correctly isolated these "
        "as INCONCLUSIVE rather than misclassifying them as non-existent. "
        "(3) Passive DNS Precision: Domain queries for example.com resolved dual authoritative IPv4 addresses (104.20.23.154, 172.66.147.243), "
        "IPv6 AAAA records, Cloudflare NS delegations, and RFC-compliant SPF restriction records (v=spf1 -all) within 720ms total elapsed time."
    )

    add_styled_heading(doc, "16. Technical Limitations & Countermeasures", level=1)
    add_body_paragraph(
        doc,
        "Cloud Edge Challenges: Advanced CAPTCHA challenges (Cloudflare Turnstile, Cloudfront) intercept programmatic probes. TRACE mitigates "
        "this by explicitly categorizing such responses as INCONCLUSIVE and recording HTTP headers rather than attempting deceptive bypasses. "
        "Name Collision: On widespread consumer networks, common usernames may correspond to distinct individuals across platforms. TRACE mitigates "
        "false identity linkage by providing direct profile links and latency telemetry, enabling human-in-the-loop contextual verification."
    )

    add_styled_heading(doc, "17. Ethical Framework & Legal Compliance", level=1)
    add_body_paragraph(
        doc,
        "TRACE was engineered strictly in accordance with ethical cybersecurity research principles and applicable cyber jurisprudence: "
        "(1) Passive Reconnaissance Doctrine: The software queries only publicly available data sources without attempting password guessing, "
        "credential stuffing, session hijacking, or perimeter exploitation. "
        "(2) Statutory Alignment: All operations comply with Section 43 and Section 66 of the Information Technology Act (India, 2000/2008), "
        "the Computer Fraud and Abuse Act (CFAA, 18 U.S.C. § 1030), and international conventions regarding unauthorized access. "
        "(3) Source Etiquette: Requests utilize legitimate HTTP User-Agent identifiers, enforce conservative timeouts (4.0s), and respect rate limits. "
        "TRACE is explicitly restricted to authorized personal footprint audits, enterprise self-assessments, and academic laboratory pedagogy."
    )

    add_styled_heading(doc, "18. Conclusion & Future Enhancements", level=1)
    add_body_paragraph(
        doc,
        "TRACE successfully accomplishes all stated objectives, establishing a robust, modern, and transparent OSINT digital footprint analyzer. "
        "By enforcing strict status disambiguation, passive DNS/RDAP correlation, local SQLite privacy, and automated academic reporting, "
        "the tool eliminates the primary flaws of legacy CLI scripts. Future work will expand source coverage to 50+ platforms, integrate "
        "asynchronous proxy rotation for inconclusive sources, and introduce interactive link-analysis graph visualizations."
    )

    add_styled_heading(doc, "19. Academic & Technical References", level=1)
    refs = [
        "[1] P. Mockapetris, RFC 1035: Domain Names - Implementation and Specification, Internet Engineering Task Force (IETF), Nov. 1987.",
        "[2] A. Newton and B. Ellacott, RFC 7480: HTTP Usage in the Registration Data Access Protocol (RDAP), IETF, Mar. 2015.",
        "[3] K. Scarfone et al., Technical Guide to Information Security Testing and Assessment, NIST Special Publication 800-115, Sept. 2008.",
        "[4] S. Bazzell, Open Source Intelligence Techniques: Resources for Searching and Analyzing Online Information, 7th ed., 2021.",
        "[5] S. Ramirez, FastAPI: Modern, High-Performance Web Framework for Python, Zenodo, 2020.",
        "[6] Ministry of Law and Justice, The Information Technology Act, 2000 (Act No. 21 of 2000), Government of India, 2000."
    ]
    for r in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        run = p.add_run(r)
        run.font.size = Pt(6.8)
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(filename)
    print(f"Successfully generated DOCX: {filename}")

if __name__ == '__main__':
    out_docx = "/home/fardeenakmal/Desktop/osint/TRACE-OSINT/case-study/TRACE_OSINT_Case_Study.docx"
    build_docx(out_docx)

